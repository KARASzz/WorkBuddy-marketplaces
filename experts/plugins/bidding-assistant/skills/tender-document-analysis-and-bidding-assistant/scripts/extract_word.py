#!/usr/bin/env python3
"""
从 Word (.docx) 提取文本和表格内容。
依赖：docx2python（优先）或 python-docx（降级）
安装：pip install docx2python python-docx
"""
import argparse
import importlib
import sys


def extract_with_docx2python(input_path: str) -> str:
    # 可选增强依赖：未安装时由调用方降级到 python-docx。
    docx2python = importlib.import_module("docx2python").docx2python
    result = docx2python(input_path)
    chunks = []
    for section in result.body:
        for table in section:
            for row in table:
                for cell in row:
                    text = "\n".join(cell).strip()
                    if text:
                        chunks.append(text)
            chunks.append("")
    return "\n".join(chunks)


def extract_with_python_docx(input_path: str) -> str:
    from docx import Document
    doc = Document(input_path)
    chunks = []
    for para in doc.paragraphs:
        if para.text.strip():
            chunks.append(para.text)
    for table in doc.tables:
        chunks.append("\n[表格内容]")
        for row in table.rows:
            chunks.append(" | ".join(cell.text.strip() for cell in row.cells))
        chunks.append("[表格结束]\n")
    return "\n".join(chunks)


def extract_word(input_path: str, output_path: str) -> None:
    text = ""
    try:
        text = extract_with_docx2python(input_path)
        print("[OK] 使用 docx2python 提取")
    except ImportError:
        pass
    except Exception as e:
        print(f"[WARN] docx2python 失败：{e}，尝试 python-docx", file=sys.stderr)

    if not text:
        try:
            text = extract_with_python_docx(input_path)
            print("[OK] 使用 python-docx 提取")
        except ImportError:
            print("[ERROR] 缺少依赖：pip install docx2python python-docx", file=sys.stderr)
            sys.exit(1)

    with open(output_path, "w", encoding="utf-8") as f:
        f.write(text)
    print(f"[OK] 提取完成，共 {len(text)} 字符 → {output_path}")


if __name__ == "__main__":
    parser = argparse.ArgumentParser(description="Word 文本提取工具")
    parser.add_argument("input", help="Word 文件路径（.docx）")
    parser.add_argument("--output", required=True, help="输出文本文件路径")
    args = parser.parse_args()
    extract_word(args.input, args.output)
