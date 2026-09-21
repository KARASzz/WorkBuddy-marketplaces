#!/usr/bin/env python3
"""
文档文本提取工具（查重专用版）。
支持 .docx / .pdf / .txt，支持单文件和批量目录两种模式。

单文件模式：
  python extract_text.py <file> --output /tmp/dedup_main.txt --structure /tmp/dedup_structure.json

批量模式：
  python extract_text.py <dir> --batch --output-dir /tmp/dedup_corpus/

依赖：python-docx（Word）、pdfplumber（PDF）
"""
import argparse
import json
import os
import re
import sys
from pathlib import Path

# 单技能打包模式：OCR/API 辅助脚本随本技能放在 scripts/ 目录内。
_LOCAL_SCRIPTS = Path(__file__).resolve().parent
sys.path.insert(0, str(_LOCAL_SCRIPTS))

MIN_CHARS_OCR = 500   # 低于此字数触发 OCR 降级


# ── 章节标题识别 ──────────────────────────────────────────────────────────────
_H1_RE = re.compile(
    r'^(第[一二三四五六七八九十百\d]+[章节部分条]'
    r'|[一二三四五六七八九十]+[、．.]'
    r'|\d+[\.、]\s*\S'
    r'|（[一二三四五六七八九十]+）)'
)
_H2_RE = re.compile(
    r'^(\d+\.\d+[\.、\s]'
    r'|[（(]\d+[）)]'
    r'|\([一二三四五六七八九十]+\))'
)

def _detect_level(text: str) -> int:
    """检测标题层级：1=一级，2=二级，0=正文。"""
    t = text.strip()
    if _H1_RE.match(t):
        return 1
    if _H2_RE.match(t):
        return 2
    return 0


# ── Word 提取 ──────────────────────────────────────────────────────────────────
def extract_word(path: str):
    try:
        from docx import Document
    except ImportError:
        sys.exit("[ERROR] 缺少依赖：pip install python-docx")

    doc      = Document(path)
    lines    = []
    sections = []
    cur_title, cur_start, cur_chars = None, 0, 0

    for para in doc.paragraphs:
        text = para.text.strip()
        if not text:
            continue
        level = _detect_level(text)
        if level == 1 or para.style.name.startswith("Heading 1"):
            if cur_title is not None:
                sections.append({
                    "title":      cur_title,
                    "char_start": cur_start,
                    "char_count": cur_chars,
                    "level":      1
                })
            cur_title  = text
            cur_start  = sum(len(l) for l in lines)
            cur_chars  = 0
        lines.append(text)
        cur_chars += len(text)

    if cur_title is not None:
        sections.append({"title": cur_title, "char_start": cur_start,
                         "char_count": cur_chars, "level": 1})

    # 表格
    for table in doc.tables:
        for row in table.rows:
            row_text = " | ".join(c.text.strip() for c in row.cells if c.text.strip())
            if row_text:
                lines.append(row_text)

    return "\n".join(lines), sections


# ── PDF 提取 ───────────────────────────────────────────────────────────────────
def extract_pdf(path: str):
    try:
        import pdfplumber
    except ImportError:
        sys.exit("[ERROR] 缺少依赖：pip install pdfplumber")

    lines, sections = [], []
    cur_title, cur_start, cur_chars = None, 0, 0

    with pdfplumber.open(path) as pdf:
        for page in pdf.pages:
            text = page.extract_text() or ""
            for line in text.splitlines():
                line = line.strip()
                if not line:
                    continue
                level = _detect_level(line)
                if level == 1:
                    if cur_title is not None:
                        sections.append({
                            "title": cur_title, "char_start": cur_start,
                            "char_count": cur_chars, "level": 1
                        })
                    cur_title = line
                    cur_start = sum(len(l) for l in lines)
                    cur_chars = 0
                lines.append(line)
                cur_chars += len(line)

            for table in (page.extract_tables() or []):
                for row in table:
                    row_text = " | ".join(cell or "" for cell in row if cell)
                    if row_text.strip():
                        lines.append(row_text.strip())

    if cur_title is not None:
        sections.append({"title": cur_title, "char_start": cur_start,
                         "char_count": cur_chars, "level": 1})

    return "\n".join(lines), sections


# ── TXT 提取 ───────────────────────────────────────────────────────────────────
def extract_txt(path: str):
    text = Path(path).read_text(encoding="utf-8", errors="replace")
    lines    = []
    sections = []
    cur_title, cur_start, cur_chars = None, 0, 0

    for line in text.splitlines():
        line = line.strip()
        if not line:
            continue
        if _detect_level(line) == 1:
            if cur_title is not None:
                sections.append({"title": cur_title, "char_start": cur_start,
                                  "char_count": cur_chars, "level": 1})
            cur_title = line
            cur_start = sum(len(l) for l in lines)
            cur_chars = 0
        lines.append(line)
        cur_chars += len(line)

    if cur_title is not None:
        sections.append({"title": cur_title, "char_start": cur_start,
                          "char_count": cur_chars, "level": 1})

    return "\n".join(lines), sections


# ── 段落切分（查重用） ─────────────────────────────────────────────────────────
_PARA_SEP = re.compile(r'\n{2,}|(?<=[。！？])\n')

def split_paragraphs(text: str, min_chars: int = 50) -> list:
    """
    将全文切分为段落列表，过滤过短段落。
    返回：[{"id": i, "text": ..., "char_count": ..., "start": ...}]
    """
    raw_paras = _PARA_SEP.split(text)
    result    = []
    offset    = 0
    pid       = 0
    for para in raw_paras:
        para = para.strip()
        if len(para) >= min_chars:
            result.append({
                "id":         pid,
                "text":       para,
                "char_count": len(para),
                "start":      text.find(para, offset)
            })
            pid += 1
        offset = text.find(para, offset) + len(para) if para in text[offset:] else offset

    # 若段落太少（<5），改用换行切分
    if len(result) < 5:
        result = []
        pid = 0
        for line in text.splitlines():
            line = line.strip()
            if len(line) >= min_chars:
                result.append({"id": pid, "text": line, "char_count": len(line), "start": 0})
                pid += 1

    return result


# ── 分发提取 ───────────────────────────────────────────────────────────────────
def auto_extract(path: str):
    """
    本地优先提取；内容 < MIN_CHARS_OCR 时自动降级到 RicheeAI OCR API。
    返回 (text: str, sections: list)。
    """
    ext = Path(path).suffix.lower()
    if ext == ".pdf":
        text, sections = extract_pdf(path)
    elif ext in (".docx", ".doc"):
        text, sections = extract_word(path)
    else:
        text, sections = extract_txt(path)

    # OCR 降级：扫描件 / 加密 PDF 本地提取内容不足时调用 API 识别
    if len(text) < MIN_CHARS_OCR:
        try:
            from ocr_extract import extract_with_ocr_fallback
            ocr_text, method = extract_with_ocr_fallback(path, min_chars=MIN_CHARS_OCR)
            if method in ("local", "ocr_api") and len(ocr_text) > len(text):
                print(
                    f"[OCR] 本地提取不足（{len(text)} 字），已通过 {method} 重新提取"
                    f"（{len(ocr_text)} 字）",
                    file=sys.stderr,
                )
                return ocr_text, []   # OCR 结果无法保留章节结构
        except Exception as e:
            print(f"[OCR] 降级失败，使用本地提取结果：{e}", file=sys.stderr)

    return text, sections


# ── 单文件模式 ─────────────────────────────────────────────────────────────────
def single_mode(args):
    path = args.input
    if not Path(path).exists():
        sys.exit(f"[ERROR] 文件不存在：{path}")

    text, sections = auto_extract(path)
    paragraphs     = split_paragraphs(text, min_chars=args.min_chars)

    if len(text) < MIN_CHARS_OCR:
        print(f"[WARN] 最终提取内容仍不足（{len(text)} 字），可能为加密文件或 OCR 失败",
              file=sys.stderr)

    Path(args.output).parent.mkdir(parents=True, exist_ok=True)
    Path(args.output).write_text(text, encoding="utf-8")

    structure = {
        "source_file":   str(Path(path).name),
        "total_chars":   len(text),
        "section_count": len(sections),
        "para_count":    len(paragraphs),
        "sections":      sections,
        "paragraphs":    paragraphs
    }
    Path(args.structure).parent.mkdir(parents=True, exist_ok=True)
    Path(args.structure).write_text(
        json.dumps(structure, ensure_ascii=False, indent=2), encoding="utf-8"
    )
    print(f"[OK] 提取完成：{len(text)} 字，{len(sections)} 个章节，"
          f"{len(paragraphs)} 个段落 → {args.output}")


# ── 批量模式 ───────────────────────────────────────────────────────────────────
def batch_mode(args):
    src_dir = Path(args.input)
    if not src_dir.is_dir():
        sys.exit(f"[ERROR] 目录不存在：{src_dir}")

    out_dir = Path(args.output_dir)
    out_dir.mkdir(parents=True, exist_ok=True)

    exts     = {".pdf", ".docx", ".doc", ".txt"}
    files    = [f for f in src_dir.iterdir() if f.suffix.lower() in exts]
    index    = []

    for f in sorted(files):
        try:
            text, sections = auto_extract(str(f))
            paragraphs     = split_paragraphs(text, min_chars=args.min_chars)
            stem           = f.stem
            out_txt        = out_dir / f"{stem}.txt"
            out_json       = out_dir / f"{stem}_structure.json"
            out_txt.write_text(text, encoding="utf-8")
            out_json.write_text(json.dumps({
                "source_file": f.name,
                "total_chars": len(text),
                "sections":    sections,
                "paragraphs":  paragraphs
            }, ensure_ascii=False, indent=2), encoding="utf-8")
            index.append({"file": f.name, "stem": stem, "chars": len(text),
                           "paras": len(paragraphs)})
            print(f"  [OK] {f.name} → {len(text)} 字 / {len(paragraphs)} 段")
        except Exception as e:
            print(f"  [WARN] {f.name} 提取失败：{e}", file=sys.stderr)
            index.append({"file": f.name, "stem": f.stem, "chars": 0,
                           "paras": 0, "error": str(e)})

    (out_dir / "_index.json").write_text(
        json.dumps(index, ensure_ascii=False, indent=2), encoding="utf-8"
    )
    print(f"[OK] 批量提取完成：{len(files)} 个文件 → {out_dir}")


# ── 入口 ───────────────────────────────────────────────────────────────────────
def main():
    parser = argparse.ArgumentParser(description="文档文本提取工具（查重专用）")
    parser.add_argument("input",        help="文件路径（单文件）或目录路径（批量）")
    parser.add_argument("--batch",      action="store_true", help="批量模式")
    parser.add_argument("--output",     default="/tmp/dedup_main.txt",     help="单文件输出路径")
    parser.add_argument("--structure",  default="/tmp/dedup_structure.json", help="结构 JSON 输出路径")
    parser.add_argument("--output-dir", default="/tmp/dedup_corpus/",      help="批量输出目录")
    parser.add_argument("--min-chars",  type=int, default=50,              help="段落最小字符数")
    args = parser.parse_args()

    if args.batch:
        batch_mode(args)
    else:
        single_mode(args)


if __name__ == "__main__":
    main()
