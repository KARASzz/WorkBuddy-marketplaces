#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
md2docx.py — 通用 Markdown → Word 装册渲染器（支持图片/表格/中文排版）。

与本项目解耦：所有内容通过"册配置 JSON"驱动，不写死任何项目路径或标题。

配置 JSON 格式：
{
  "out": "/path/输出.docx",
  "title": "封面主标题",
  "subtitle": "副标题",
  "cover_lines": ["采购方案编号：xxx", "竞价单位：xxx"],
  "cover_warning": "【看效果版 · 非递交版】",
  "cover_notes": ["提示1", "提示2"],
  "sections": [
    {"heading": "第一部分  技术册", "file": "/path/技术册.md"},
    {"heading": "第二部分  商务册", "file": "/path/商务册.md"}
  ],
  "page": {"width": 8.27, "height": 11.69, "margin": 1.0},
  "img_width_inches": 6.2
}

支持的 Markdown 语法：
  # / ## / ### 标题
  | 表格 |（自动套 Table Grid）
  ![图注](绝对路径或相对md文件的路径)  ← 图片，按页宽缩放居中
  > 引用（含"待/阻断/占位/注意"等词自动转为红色提示样式）
  - / * 列表，1. 有序列表
  ---  分隔线
  ```代码块```
  YAML 头（--- 包裹）自动跳过

用法：
  python3 md2docx.py --config config.json
"""
import argparse
import json
import os
import re
import sys

from docx import Document
from docx.shared import Pt, RGBColor, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

HINT_WORDS = ['待', '阻断', 'pending', '占位', '声明', '重要', '注意', '提示', '版本', '须', '缺失']


def set_cn_font(run, name="宋体", size=10.5, bold=False, color=None):
    run.font.name = name
    run.font.size = Pt(size)
    run.font.bold = bold
    if color:
        run.font.color.rgb = RGBColor(*color)
    rPr = run._element.get_or_add_rPr()
    rFonts = rPr.find(qn('w:rFonts'))
    if rFonts is None:
        rFonts = OxmlElement('w:rFonts')
        rPr.append(rFonts)
    rFonts.set(qn('w:eastAsia'), name)


def add_para(doc, text, size=10.5, bold=False, color=None, align=None, space_after=4, indent=None):
    p = doc.add_paragraph()
    if align is not None:
        p.alignment = align
    p.paragraph_format.space_after = Pt(space_after)
    if indent is not None:
        p.paragraph_format.left_indent = Inches(indent)
    run = p.add_run(text)
    set_cn_font(run, size=size, bold=bold, color=color)
    return p


def add_heading(doc, text, level=1):
    sizes = {1: 16, 2: 13, 3: 11.5}
    colors = {1: (0x1F, 0x3B, 0x73), 2: (0x2E, 0x5B, 0x9F), 3: (0x33, 0x33, 0x33)}
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(10 if level == 1 else 6)
    p.paragraph_format.space_after = Pt(4)
    run = p.add_run(text)
    set_cn_font(run, size=sizes.get(level, 11), bold=True, color=colors.get(level))
    return p


def add_hint(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(3)
    run = p.add_run(text)
    set_cn_font(run, size=9.5, color=(0xB0, 0x30, 0x20))
    return p


def add_picture(doc, img_path, caption=None, max_width_inches=6.2):
    if not os.path.exists(img_path):
        add_hint(doc, f"[图片缺失: {img_path}]")
        return False
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(2)
    try:
        p.add_run().add_picture(img_path, width=Inches(max_width_inches))
    except Exception as e:
        add_hint(doc, f"[图片插入失败 {os.path.basename(img_path)}: {e}]")
        return False
    if caption:
        cp = doc.add_paragraph()
        cp.alignment = WD_ALIGN_PARAGRAPH.CENTER
        cp.paragraph_format.space_after = Pt(6)
        set_cn_font(cp.add_run(caption), size=8.5, color=(0x66, 0x66, 0x66))
    return True


def add_table(doc, header, rows):
    ncols = len(header)
    tbl = doc.add_table(rows=1, cols=ncols)
    tbl.style = 'Table Grid'
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl.autofit = False
    for i, h in enumerate(header):
        c = tbl.rows[0].cells[i]
        c.text = ""
        set_cn_font(c.paragraphs[0].add_run(h), size=9, bold=True)
    for row in rows:
        cells = tbl.add_row().cells
        for i in range(ncols):
            val = row[i] if i < len(row) else ""
            cells[i].text = ""
            set_cn_font(cells[i].paragraphs[0].add_run(val), size=9)
    return tbl


def parse_md_table(lines, start):
    i = start
    while i < len(lines) and '|' not in lines[i]:
        i += 1
    if i >= len(lines):
        return None, None, start
    header = [c.strip() for c in lines[i].strip().strip('|').split('|')]
    i += 1
    if i < len(lines) and re.match(r'^\s*\|?[\s:|-]+\|?\s*$', lines[i]):
        i += 1
    rows = []
    while i < len(lines) and lines[i].strip().startswith('|'):
        rows.append([c.strip() for c in lines[i].strip().strip('|').split('|')])
        i += 1
    return header, rows, i


def render_md_file(doc, path, skip_yaml=True, img_width=6.2):
    """渲染单个 Markdown 文件。返回 (图片数, 缺失图片数)"""
    base_dir = os.path.dirname(path)
    with open(path, encoding='utf-8') as f:
        lines = f.read().split('\n')

    i, in_code = 0, False
    n_img = n_missing = 0
    if skip_yaml and lines and lines[0].strip() == '---':
        i = 1
        while i < len(lines) and lines[i].strip() != '---':
            i += 1
        i += 1

    while i < len(lines):
        s = lines[i].rstrip().strip()
        if s.startswith('```'):
            in_code = not in_code
            i += 1
            continue
        if in_code:
            add_para(doc, lines[i].rstrip(), size=8.5, color=(0x44, 0x44, 0x44))
            i += 1
            continue
        if not s:
            i += 1
            continue
        # 图片（整行 ![...](...)）
        m = re.match(r'^!\[(.*?)\]\((.+?)\)\s*$', s)
        if m:
            caption, rel = m.group(1).strip(), m.group(2).strip()
            ip = rel if os.path.isabs(rel) else os.path.join(base_dir, rel)
            if add_picture(doc, ip, caption=caption, max_width_inches=img_width):
                n_img += 1
            else:
                n_missing += 1
            i += 1
            continue
        # 标题
        if s.startswith('#'):
            lv = min(len(s) - len(s.lstrip('#')), 3)
            add_heading(doc, s.lstrip('#').strip(), lv)
            i += 1
            continue
        # 表格
        if s.startswith('|'):
            header, rows, ni = parse_md_table(lines, i)
            if header:
                add_table(doc, header, rows)
                i = ni
                continue
        # 引用
        if s.startswith('>'):
            txt = s.lstrip('>').strip()
            if any(k in txt for k in HINT_WORDS):
                add_hint(doc, txt)
            else:
                add_para(doc, txt, size=9.5, color=(0x55, 0x55, 0x55))
            i += 1
            continue
        if re.match(r'^[-*_]{3,}$', s):
            i += 1
            continue
        if s.startswith('- ') or s.startswith('* '):
            add_para(doc, s[2:], size=10, indent=0.25)
            i += 1
            continue
        if re.match(r'^\d+\.\s', s):
            add_para(doc, s, size=10, indent=0.25)
            i += 1
            continue
        add_para(doc, s, size=10)
        i += 1
    return n_img, n_missing


def build(cfg):
    doc = Document()
    pg = cfg.get('page', {})
    for sec in doc.sections:
        sec.page_width = Inches(pg.get('width', 8.27))
        sec.page_height = Inches(pg.get('height', 11.69))
        m = pg.get('margin', 1.0)
        sec.top_margin = sec.bottom_margin = Inches(m)
        sec.left_margin = sec.right_margin = Inches(m)

    # 封面
    add_para(doc, "", size=12)
    if cfg.get('title'):
        add_para(doc, cfg['title'], size=20, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER)
    if cfg.get('subtitle'):
        add_para(doc, cfg['subtitle'], size=28, bold=True,
                 align=WD_ALIGN_PARAGRAPH.CENTER, space_after=12)
    for ln in cfg.get('cover_lines', []):
        add_para(doc, ln, size=12, align=WD_ALIGN_PARAGRAPH.CENTER)
    add_para(doc, "", size=12)
    if cfg.get('cover_warning'):
        add_para(doc, cfg['cover_warning'], size=13, bold=True,
                 color=(0xB0, 0x30, 0x20), align=WD_ALIGN_PARAGRAPH.CENTER)
    for n in cfg.get('cover_notes', []):
        add_hint(doc, n)
    doc.add_page_break()

    img_w = cfg.get('img_width_inches', 6.2)
    total_img = total_missing = 0
    for s in cfg.get('sections', []):
        if s.get('heading'):
            add_heading(doc, s['heading'], 1)
        if s.get('file'):
            if not os.path.exists(s['file']):
                add_hint(doc, f"[文件缺失: {s['file']}]")
                continue
            n, miss = render_md_file(doc, s['file'], img_width=img_w)
            total_img += n
            total_missing += miss
            print(f"  {os.path.basename(s['file'])}: 图片 {n} 张，缺失 {miss} 张")
        if s.get('page_break', True):
            doc.add_page_break()

    out = cfg['out']
    os.makedirs(os.path.dirname(out) or '.', exist_ok=True)
    doc.save(out)
    print(f"\n已生成: {out}")
    if os.path.exists(out):
        print(f"  大小: {os.path.getsize(out)/1024/1024:.1f} MB")
    print(f"  图片合计: {total_img} 张，缺失 {total_missing} 张")


def main():
    ap = argparse.ArgumentParser(description='通用 Markdown → Word 装册渲染器')
    ap.add_argument('--config', required=True, help='册配置 JSON')
    args = ap.parse_args()
    with open(args.config, encoding='utf-8') as f:
        cfg = json.load(f)
    build(cfg)


if __name__ == '__main__':
    main()
